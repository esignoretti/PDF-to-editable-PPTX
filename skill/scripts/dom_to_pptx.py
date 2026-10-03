"""Build an editable PPTX from a DOM spec produced by extract_dom.py.

This is the HTML fast path's second half: no pixel measurement, no font
calibration — positions, colours, sizes and tracking come straight from the
browser's computed style.

    dom_to_pptx.py --spec spec.json --out deck.pptx
    dom_to_pptx.py --html deck.html --out deck.pptx --slide 2   # extract + build

The output is deliberately conservative: it emits native shapes and real text
boxes and skips anything it cannot express (see extract_dom's warnings). Always
verify with the raster diff — a deck that leans on transforms, box-shadows or
radial gradients will want the raster route instead.
"""
import argparse
import json
import math
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pptx_helpers as H  # noqa: E402

ALIGN = {"left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
DEFAULT_WIDTH_IN = 13.333
DEFAULT_HEIGHT_IN = 7.5


def _hex(value, fallback="000000"):
    v = (value or fallback).lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    return v[:6].upper()


def _grad_angle(css_deg):
    """CSS gradient angle (0 = up, clockwise) -> DrawingML (0 = east, clockwise)."""
    rad = math.radians(css_deg)
    dx, dy = math.sin(rad), -math.cos(rad)
    deg = math.degrees(math.atan2(dy, dx)) % 360.0
    return int(round(deg * 60000)) % 21600000


def _font_for(family, weight):
    return family or "Arial", int(weight or 400) >= 600


def build(spec, out, width_in=DEFAULT_WIDTH_IN, height_in=None, assets_dir=None,
          dry_run=False):
    sw = float(spec["slide"]["w"])
    sh = float(spec["slide"]["h"])
    if height_in is None:
        height_in = DEFAULT_HEIGHT_IN if abs(sw / sh - 16 / 9) < 0.02 else width_in * sh / sw

    pt_per_px = width_in * 72.0 / sw
    in_per_px = width_in / sw

    def px(v, total):
        return v / total * 100.0

    prs = Presentation()
    prs.slide_width = Emu(int(round(width_in * 914400)))
    prs.slide_height = Emu(int(round(height_in * 914400)))
    H.SW, H.SH = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = spec["slide"].get("background")
    bgg = spec["slide"].get("gradient")
    if (bgg and bgg.get("type") == "linear") or (bg and bg.get("a", 1) > 0.01):
        rect = slide.shapes.add_shape(1, 0, 0, Emu(H.SW), Emu(H.SH))
        rect.shadow.inherit = False
        if bgg and bgg.get("type") == "linear":
            H.set_fill(rect, grad=[(s["pos"], _hex(s["hex"]), round(s.get("a", 1) * 100))
                                   for s in bgg["stops"]], ang=_grad_angle(bgg["css_angle"]))
        else:
            H.set_fill(rect, _hex(bg["hex"]), alpha=round(bg.get("a", 1) * 100))
        H.no_line(rect)

    stats = {"box": 0, "text": 0, "svg": 0, "image": 0, "skipped": 0}
    for node in spec["nodes"]:
        b = node.get("box") or {}
        if not b or b.get("w", 0) <= 0.1 or b.get("h", 0) <= 0.1:
            stats["skipped"] += 1
            continue
        x, y = px(b["x"], sw), px(b["y"], sh)
        w, h = px(b["w"], sw), px(b["h"], sh)

        if node["kind"] == "box":
            grad = node.get("gradient")
            grad_stops = None
            if grad and grad.get("type") == "linear":
                grad_stops = [(s["pos"], _hex(s["hex"]), round(s.get("a", 1) * 100))
                              for s in grad["stops"]]
            fill = node.get("fill")
            border = node.get("border")
            H.roundrect(
                slide, x, y, x + w, y + h,
                radius_in=node.get("radius_px", 0) * in_per_px,
                grad=grad_stops,
                ang=_grad_angle(grad["css_angle"]) if grad_stops else 5400000,
                fill=None if grad_stops else (_hex(fill["hex"]) if fill else None),
                alpha=None if grad_stops or not fill else round(fill.get("a", 1) * 100),
                line=_hex(border["hex"]) if border else None,
                lw=border["w"] * pt_per_px if border else 0.75,
            )
            stats["box"] += 1

        elif node["kind"] == "text":
            runs = []
            for r in node["runs"]:
                st = r["style"]
                text = r["text"]
                if st.get("transform") == "uppercase":
                    text = text.upper()
                elif st.get("transform") == "lowercase":
                    text = text.lower()
                fam, bold = _font_for(st.get("family"), st.get("weight"))
                runs.append((text, fam, st["size_px"] * pt_per_px, bold,
                             _hex(st.get("color")), st.get("letter_spacing_px", 0) * pt_per_px))
            if not runs:
                stats["skipped"] += 1
                continue
            line_h_px = node.get("line_height_px", 0) or 0
            single_line = bool(line_h_px) and b["h"] <= line_h_px * 1.5
            wrap = (not single_line) and node.get("white_space", "normal") not in ("nowrap", "pre")
            H.text(slide, x, y, w, runs, align=ALIGN.get(node.get("align", "left"), PP_ALIGN.LEFT),
                   h=h, top=y, anchor=MSO_ANCHOR.TOP,
                   line_spacing_pt=line_h_px * pt_per_px or None,
                   wrap=wrap)
            stats["text"] += 1

        elif node["kind"] in ("svg", "image"):
            asset = node.get("asset")
            if not asset:
                stats["skipped"] += 1
                continue
            path = asset if os.path.isabs(asset) else os.path.join(assets_dir or ".", os.path.basename(asset))
            if not os.path.exists(path):
                stats["skipped"] += 1
                continue
            H.picture(slide, path, x, y, w, h)
            stats[node["kind"]] += 1

    if not dry_run:
        prs.save(out)
    return {"out": out, "slide_in": [round(width_in, 3), round(height_in, 3)],
            "pt_per_px": round(pt_per_px, 4), "nodes": stats,
            "warnings": spec.get("warnings", [])}


def _main():
    ap = argparse.ArgumentParser(description="Build an editable PPTX from a DOM spec.")
    ap.add_argument("--spec", help="spec JSON from extract_dom.py")
    ap.add_argument("--html", help="HTML source: extract then build in one step")
    ap.add_argument("--out", "-o", default="deck.pptx")
    ap.add_argument("--assets", default=None)
    ap.add_argument("--slide", type=int, default=0)
    ap.add_argument("--selector", default=".slide")
    ap.add_argument("--width", type=int, default=1920)
    ap.add_argument("--height", type=int, default=1080)
    ap.add_argument("--slide-width-in", type=float, default=DEFAULT_WIDTH_IN)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not args.spec and not args.html:
        ap.error("pass --spec or --html")

    if args.html:
        import extract_dom
        spec_path = os.path.splitext(args.out)[0] + ".spec.json"
        assets = args.assets or os.path.splitext(args.out)[0] + ".assets"
        spec = extract_dom.extract(args.html, spec_path, assets, args.selector,
                                   args.slide, args.width, args.height)
        assets_dir = assets
    else:
        with open(args.spec) as fh:
            spec = json.load(fh)
        assets_dir = args.assets or os.path.dirname(os.path.abspath(args.spec))

    report = build(spec, args.out, args.slide_width_in, None, assets_dir, args.dry_run)
    print(json.dumps(report, indent=1))
    if report["warnings"]:
        print("\nThe source used features this path punts on. Check the output, or use the "
              "raster route (render_html -> measure) for those slides.", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(_main())
